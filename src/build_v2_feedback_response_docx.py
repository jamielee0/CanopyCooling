#!/usr/bin/env python3
"""Build the formal, nonthermal response to all 37 reviewer feedback points."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Iterable
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/v2/hitl/FEEDBACK_CLOSEOUT"
MATRIX = OUT / "feedback_compliance_matrix.csv"
CHECKS = OUT / "checks.json"
G4 = OUT / "g4_checks.json"
G5 = OUT / "g5_checks.json"
G6 = OUT / "g6_independent_verification.json"
G7 = OUT / "g7_simulation_contract.json"
G8 = OUT / "g8_decision.json"
OUTPUT = OUT / "Urban_Tree_Cooling_Feedback_Response.docx"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

BLACK = "111111"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5F6670"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D7DBE2"
CALLOUT = "F4F6F9"
APPROVAL_FILL = "EAF3ED"  # Named approval-status override.
APPROVAL_INK = "1B5E3C"
CAUTION_FILL = "FFF5D6"  # Named scientific-scope warning override.
CAUTION_INK = "7A5A00"
NEGATIVE_FILL = "FDECEC"
NEGATIVE_INK = "8C1D18"


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def read_matrix() -> list[dict[str, str]]:
    with MATRIX.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def set_run(
    run,
    *,
    size: float = 11,
    color: str = BLACK,
    bold: bool | None = None,
    italic: bool | None = None,
    font: str = "Calibri",
) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in CELL_MARGIN_DXA.items():
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa: Iterable[int]) -> None:
    widths = list(widths_dxa)
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}, received {widths}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        prevent_row_split(row)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def set_table_paragraph(paragraph, *, size: float = 9, color: str = BLACK, bold: bool = False) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    for run in paragraph.runs:
        set_run(run, size=size, color=color, bold=bold)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
    *,
    body_size: float = 9,
    status_column: int | None = None,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr()
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.add_run(text)
        set_table_paragraph(paragraph, size=9, color=DARK_BLUE, bold=True)
    set_repeat_table_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, values)):
            paragraph = cell.paragraphs[0]
            paragraph.add_run(value)
            color = BLACK
            if status_column is not None and index == status_column:
                if value.startswith("Addressed — negative"):
                    set_cell_shading(cell, NEGATIVE_FILL)
                    color = NEGATIVE_INK
                elif value.startswith("Partial"):
                    set_cell_shading(cell, CAUTION_FILL)
                    color = CAUTION_INK
                else:
                    set_cell_shading(cell, APPROVAL_FILL)
                    color = APPROVAL_INK
            set_table_paragraph(paragraph, size=body_size, color=color, bold=index == 0)
    set_table_geometry(table, widths_dxa)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(4)


def add_heading(doc: Document, text: str, level: int, *, page_break_before: bool = False) -> None:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break_before


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run(lead, bold=True)
        body = paragraph.add_run(text[len(bold_lead) :])
        set_run(body)
    else:
        run = paragraph.add_run(text)
        set_run(run)


def add_metadata(doc: Document, label: str, value: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    label_run = paragraph.add_run(f"{label}: ")
    set_run(label_run, bold=True)
    value_run = paragraph.add_run(value)
    set_run(value_run)


def add_bottom_rule(paragraph, color: str = BLUE, size: str = "16") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run(run, size=9, color=MUTED)
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (fld_char, instr, separate, result, end):
        run._r.append(node)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    settings = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in settings.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
    citation = doc.styles.add_style("Table Citation", WD_STYLE_TYPE.PARAGRAPH)
    citation.base_style = normal
    citation.font.name = "Calibri"
    citation.font.size = Pt(9)
    citation.font.color.rgb = RGBColor.from_string(MUTED)
    citation.paragraph_format.space_before = Pt(4)
    citation.paragraph_format.space_after = Pt(4)


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True
    for header_part in (section.header, section.first_page_header):
        header_part.is_linked_to_previous = False
        header = header_part.paragraphs[0]
        header.text = "URBAN TREE COOLING  |  FEEDBACK CLOSEOUT"
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header.paragraph_format.space_after = Pt(0)
        set_run(header.runs[0], size=8.5, color=MUTED, bold=True)
    for footer_part in (section.footer, section.first_page_footer):
        footer_part.is_linked_to_previous = False
        footer = footer_part.paragraphs[0]
        add_page_number(footer)


def configure_document(doc: Document) -> None:
    doc.settings.odd_and_even_pages_header_footer = False
    configure_section(doc.sections[0])
    settings = doc.settings._element
    compat = settings.find(qn("w:compat"))
    if compat is not None:
        fe_layout = compat.find(qn("w:useFELayout"))
        if fe_layout is not None:
            compat.remove(fe_layout)
        for option in list(compat.findall(qn("w:compatSetting"))):
            if option.get(qn("w:name")) == "doNotFlipMirrorIndents":
                compat.remove(option)
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def start_new_page_section(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section)


def add_callout(doc: Document, heading: str, text: str, *, approval: bool) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    fill = APPROVAL_FILL if approval else CAUTION_FILL
    ink = APPROVAL_INK if approval else CAUTION_INK
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    heading_run = paragraph.add_run(f"{heading}\n")
    set_run(heading_run, size=11, color=ink, bold=True)
    body_run = paragraph.add_run(text)
    set_run(body_run, size=10.5, color=BLACK)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def display_status(value: str) -> str:
    return {
        "ADDRESSED": "Addressed",
        "ADDRESSED_NEGATIVE_RESULT": "Addressed — negative result",
        "PARTIAL_BLOCKED_REPORTED": "Partial — blocked and reported",
        "ADDRESSED_EQUIVALENT_LABEL": "Addressed — equivalent control",
        "ADDRESSED_BY_FAIL_CLOSED_PREFLIGHT": "Addressed — fail-closed preflight",
    }[value]


def add_response_item(doc: Document, item: dict[str, str]) -> None:
    status = display_status(item["closeout_status"])
    if status.startswith("Addressed — negative"):
        status_color = NEGATIVE_INK
    elif status.startswith("Partial"):
        status_color = CAUTION_INK
    else:
        status_color = APPROVAL_INK
    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(5)
    label.paragraph_format.space_after = Pt(2)
    id_run = label.add_run(item["feedback_id"])
    set_run(id_run, size=10, color=DARK_BLUE, bold=True)
    sep_run = label.add_run("  |  ")
    set_run(sep_run, size=10, color=MUTED)
    status_run = label.add_run(status)
    set_run(status_run, size=10, color=status_color, bold=True)

    feedback = doc.add_paragraph()
    feedback.paragraph_format.space_after = Pt(2)
    feedback.paragraph_format.line_spacing = 1.05
    lead = feedback.add_run("Feedback: ")
    set_run(lead, size=9.5, bold=True)
    body = feedback.add_run(item["feedback_point"])
    set_run(body, size=9.5)

    response = doc.add_paragraph()
    response.paragraph_format.space_after = Pt(6)
    response.paragraph_format.line_spacing = 1.05
    lead = response.add_run("Response: ")
    set_run(lead, size=9.5, bold=True)
    body = response.add_run(
        f"{item['closeout_evidence']} Evidence: {item['closeout_gate']}."
    )
    set_run(body, size=9.5)
    add_bottom_rule(response, color=MID_GRAY, size="4")


def build() -> Path:
    matrix = read_matrix()
    checks = read_json(CHECKS)
    g4 = read_json(G4)
    g5 = read_json(G5)
    g6 = read_json(G6)
    g7 = read_json(G7)
    g8 = read_json(G8)
    if len(matrix) != 37 or not checks["all_feedback_points_have_dispositions"]:
        raise RuntimeError("Cannot author response: the 37-point compliance census is incomplete")
    if not g6["all_passed"]:
        raise RuntimeError("Cannot author response: independent verification did not pass")

    doc = Document()
    configure_styles(doc)
    configure_document(doc)
    doc.core_properties.title = "Urban Tree Cooling Study — Response to Feedback"
    doc.core_properties.subject = "Audited disposition of 37 reviewer feedback points"
    doc.core_properties.author = "Urban Tree Cooling project"
    doc.core_properties.comments = (
        "rfi_response preset; memo_masthead; named approval, caution, and negative-status fills"
    )

    top = doc.add_paragraph()
    top.paragraph_format.space_before = Pt(12)
    top.paragraph_format.space_after = Pt(3)
    run = top.add_run("FEEDBACK RESPONSE AND CLOSEOUT")
    set_run(run, size=23, color=BLACK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Urban Tree Cooling Study — nonthermal design review")
    set_run(run, size=14, color=MUTED)
    add_metadata(doc, "Review date", "August 16, 2026")
    add_metadata(doc, "Decision basis", "D0083 — user-authorized nonthermal feedback closeout")
    add_metadata(doc, "Source reviewed", "Urban_Tree_Cooling_feedbackdocx (1).docx (3 pages)")
    add_metadata(doc, "Coverage", "37 of 37 substantive feedback points assigned explicit dispositions")
    add_metadata(doc, "Review status", "Approved as a complete feedback-response packet")
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(10)
    add_bottom_rule(rule)

    add_callout(
        doc,
        "APPROVED — FEEDBACK CLOSEOUT PACKET",
        "The response matrix, nonthermal diagnostics, and independent checks are complete. "
        "Every feedback point has an evidence-backed disposition.",
        approval=True,
    )
    add_callout(
        doc,
        "SCIENTIFIC SCOPE REMAINS CLOSED",
        "This approval does not approve a pooled high-demand effect, an AM–PM comparison, "
        "thermal/LST inspection, holdout selection, Task-2 outcome work, or access to the 2026 "
        "science record. Those routes failed the frozen support rules or remain sealed.",
        approval=False,
    )

    add_heading(doc, "Executive decision", 1)
    add_body(
        doc,
        "The frozen decision rule selects COMPARATIVE_CASE_STUDY_PROPOSAL_ONLY. Denver–Aurora "
        "is the sole candidate at the widest selectable 25° diagnostic threshold. The result is "
        "a proposal direction, not an outcome estimate and not permission to open thermal data.",
    )
    add_table(
        doc,
        ["Question", "Audited result", "Authorization"],
        [
            ["Pooled high-demand wet–dry contrast", "Only 1 of 4 city-windows passes; Denver–Aurora only", "Not authorized"],
            ["Empirically matched AM–PM comparison", "0 morning passes and 0 cross-stratum pairs at 25°", "Not authorized"],
            ["Feedback-compliant simulation", "All three profiles fail empirical identifiability preflight", "Not run"],
            ["Comparative case-study proposal", "Denver–Aurora meets within-city support at 25°", "Proposal only"],
            ["Independent verification", "10 of 10 checks passed", "Accepted"],
        ],
        [3000, 4200, 2160],
        body_size=9.2,
    )

    start_new_page_section(doc)
    add_heading(doc, "Disposition summary", 1)
    counts = checks["feedback_status_counts"]
    add_table(
        doc,
        ["Disposition", "Count", "Meaning"],
        [
            ["Addressed", str(counts.get("ADDRESSED", 0)), "Requested evidence, rule, or clarification delivered"],
            ["Addressed — negative result", str(counts.get("ADDRESSED_NEGATIVE_RESULT", 0)), "Requested audit completed; support criterion failed"],
            ["Addressed — equivalent control", str(counts.get("ADDRESSED_EQUIVALENT_LABEL", 0)), "Existing quarantine control is substantively equivalent"],
            ["Addressed — fail-closed preflight", str(counts.get("ADDRESSED_BY_FAIL_CLOSED_PREFLIGHT", 0)), "Simulation contract frozen; invalid run refused"],
            ["Partial — blocked and reported", str(counts.get("PARTIAL_BLOCKED_REPORTED", 0)), "Available evidence reported; remaining field depends on an unselected design"],
            ["Total", str(len(matrix)), "Complete feedback census"],
        ],
        [3000, 1000, 5360],
        body_size=9.2,
    )

    add_heading(doc, "Gate-by-gate response", 1)
    add_heading(doc, "G0–G3 — protocol, reconciliation, wetness, and sampling design", 2)
    add_body(
        doc,
        "The outcome-blind amendment, attrition ledger, pass-equivalent interpretation, raw-count "
        "reporting, 30/60-day precipitation and P−ET0 support, days-since-rain metric, city provenance, "
        "phenology, season, geometry, and threshold audits were completed. The prospective G3 selector "
        "returned REVISE_REQUIRED, so no design was selected and no outcome-bearing branch was opened.",
    )
    add_heading(doc, "G4 — focused high-demand audit", 2)
    t25 = g4["thresholds"]["25"]
    add_body(
        doc,
        f"At 25°, 52 physical passes are available across the four primary city-windows. Only "
        f"Denver–Aurora passes the within-city wet/dry count plus positive absolute-VPD and seasonal "
        f"overlap rules. Its share of the observed passes is {t25['maximum_city_share']:.2%}, above "
        "the 40% concentration cap. The four-city pooled contrast is therefore not authorized.",
    )
    add_heading(doc, "G5 — time-of-day matching", 2)
    add_body(
        doc,
        "The complete design has no late-morning support in any primary city-window at any selectable "
        "threshold. At 25° it has 0 late-morning passes, 18 late-afternoon passes, and 0 possible "
        "cross-stratum pairs even before solar-elevation, weather, seasonal, or relative-azimuth "
        "calipers. The AM–PM comparison is not authorized.",
    )
    add_heading(doc, "G6 — independent verification", 2)
    add_body(
        doc,
        "A standalone verifier that imports no project analysis modules passed all 10 checks: source "
        "hash, 2,968-row identity and deduplication, the 2,968→219→102 attrition counts, the "
        "438-link/413-asset/219-pass weather identity, VPD units, the published NREL solar-position "
        "known answer, prior-day window exclusion, cloud/fill semantics, P−ET0 sign and units, and the "
        "Gate-3 geometry/weather/support clean room.",
    )
    start_new_page_section(doc)
    add_heading(doc, "G7 — simulation contract and preflight", 2)
    historical = g7["historical_d0056_time_simulation"]
    add_body(
        doc,
        "The three requested profiles were preregistered with 0.75/1.50/2.25 K effects, explicit "
        "pass/matched-set/row residual components (0.8/0.5/1.5 K), clustering, pass-level measurement "
        "error, alpha 0.05, target power 0.80, replicate counts, seed, and an MDE rule. No new Monte "
        "Carlo run was made because the empirical profiles are not identifiable. The historical D0056 "
        f"run used {historical['replicates']} replicates and produced power {historical['power_at_primary_effect']:.2f} "
        "at 1.5 K, but its same-distribution archive scope cannot substitute for the requested matched design.",
    )
    add_heading(doc, "G8 — final rule", 2)
    add_body(
        doc,
        "Proceed only with a nonthermal Denver–Aurora comparative case-study proposal. Reject the "
        "pooled high-demand claim, reject the AM–PM outcome analysis, do not manufacture power for "
        "non-identifiable profiles, and keep all sealed outcome branches closed.",
    )

    add_heading(doc, "Point-by-point response matrix", 1)
    add_body(
        doc,
        "The wording below follows the 37-point traceability census derived from the three-page source "
        "document. “Negative result” means the requested check was performed and did not meet its "
        "frozen support rule; it is not an omitted response.",
    )
    by_id = {item["feedback_id"]: item for item in matrix}
    matrix_page_chunks = [
        [("Source page 1", range(1, 6))],
        [("Source page 1 — continued", range(6, 15))],
        [
            ("Source page 1 — final", range(15, 17)),
            ("Source page 2", range(17, 24)),
        ],
        [
            ("Source page 2 — continued", range(24, 28)),
            ("Source page 3", range(28, 33)),
        ],
        [("Source page 3 — continued", range(33, 38))],
    ]
    for chunk_index, sections in enumerate(matrix_page_chunks):
        for label, identifiers in sections:
            add_heading(doc, label, 2)
            for identifier in identifiers:
                add_response_item(doc, by_id[f"F{identifier:02d}"])
        if chunk_index < len(matrix_page_chunks) - 1:
            start_new_page_section(doc)

    add_heading(doc, "Audit boundaries and evidence", 1)
    add_body(
        doc,
        "Source document SHA-256: 4d369a5337a20b2969d4c1ccabb528c56fae9eb19061c72513d301658dc79333.",
    )
    add_body(
        doc,
        "Primary machine-readable evidence: feedback_compliance_matrix.csv; g4_checks.json; "
        "g5_checks.json; g6_independent_verification.json; g7_simulation_contract.json; "
        "g8_decision.json; checks.json; and source_bindings.json in docs/v2/hitl/FEEDBACK_CLOSEOUT.",
    )
    add_body(
        doc,
        "Solar known-answer reference: NREL, Solar Position Algorithm for Solar Radiation Applications, "
        "https://www.nrel.gov/docs/fy08osti/34302.pdf.",
    )
    add_callout(
        doc,
        "FINAL APPROVAL",
        "Approved as a complete and independently checked response to the reviewer feedback. The only "
        "approved continuation is preparation of a nonthermal Denver–Aurora comparative case-study "
        "proposal. Outcome-bearing analyses remain unapproved and sealed.",
        approval=True,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    audit_docx(OUTPUT, expected_tables=len(doc.tables))
    return OUTPUT


def audit_docx(path: Path, *, expected_tables: int) -> None:
    doc = Document(path)
    expected = {
        "page_width": 12240,
        "page_height": 15840,
        "top_margin": 1440,
        "right_margin": 1440,
        "bottom_margin": 1440,
        "left_margin": 1440,
    }
    for section_index, section in enumerate(doc.sections, start=1):
        for name, value in expected.items():
            observed = int(getattr(section, name).twips)
            if observed != value:
                raise AssertionError(
                    f"section {section_index} {name}: expected {value}, observed {observed}"
                )
    if len(doc.tables) != expected_tables:
        raise AssertionError("Table count changed after save")
    for index, table in enumerate(doc.tables, start=1):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        widths = [int(node.get(qn("w:w"))) for node in table._tbl.tblGrid]
        if tbl_w is None or int(tbl_w.get(qn("w:w"))) != CONTENT_WIDTH_DXA:
            raise AssertionError(f"Table {index} has invalid width")
        if tbl_ind is None or int(tbl_ind.get(qn("w:w"))) != TABLE_INDENT_DXA:
            raise AssertionError(f"Table {index} has invalid indent")
        if sum(widths) != CONTENT_WIDTH_DXA:
            raise AssertionError(f"Table {index} grid does not sum to {CONTENT_WIDTH_DXA}")
        for row in table.rows:
            cell_widths = [
                int(cell._tc.get_or_add_tcPr().find(qn("w:tcW")).get(qn("w:w")))
                for cell in row.cells
            ]
            if cell_widths != widths:
                raise AssertionError(f"Table {index} cell/grid widths disagree")
    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    for feedback_id in [f"F{index:02d}" for index in range(1, 38)]:
        if feedback_id not in document_xml:
            raise AssertionError(f"Missing feedback row {feedback_id}")


def main() -> int:
    path = build()
    print(f"PASS build_v2_feedback_response_docx: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
